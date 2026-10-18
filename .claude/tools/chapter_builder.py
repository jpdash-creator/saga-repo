# ═══════════════════════════════════════════════════════════════════════════
# STARWARD SAGA — Chapter Builder (condensed Chapter .docx + .md writer)
# Authorised Novel-Output Script #6 (guardrails/tools/guardrail-authorized-scripts-only.md),
# added by author ruling on 2026-10-07 together with the /chapter-builder command, which replaces /export-chapter.
#
# What it does: writes ONE Chapter as a coupled pair, in the Novel-Part folder:
#   projects/novel-starward-saga/phase-N/section-N/part-N/Chapter NNN - <Chapter Name>.docx
#   projects/novel-starward-saga/phase-N/section-N/part-N/Chapter NNN - <Chapter Name>.md
# from one UTF-8 text file: the confirmed condensed Chapter text. The text is written by the Area Maker
# from the Chapter Spec and the Chapter's written Areas, staged in .claude/temp/, and reviewed and confirmed
# by the author; this script writes no prose, no date and no epigraph (it only typesets what it is handed).
#
# The Chapter text file, in order:
#   *Year 1247, ...*        the Chapter's date line, the first paragraph (Part 41 Area 5)
#   narration paragraphs    blank-line separated; *"..."* dialogue; ***Move*** bold-italic; **text** bold
#   ---                     a scene break (an Area boundary). Nothing else marks an Area boundary: no Area title,
#                           no reference code. md: a "---" line. docx: a centred "*   *   *" line.
#   *Year 1247, ...*        directly after a "---": the date line of that scene, present only when the
#                           date differs from the date before it (the same date is not repeated).
# Where a scene carries its own date line the docx gives that scene's first paragraph the drop cap and small
# caps, as every Area opening has; a scene with no date line starts with no indent.
#
# Look: the same typeset as the Area files (6x9 trim, EB Garamond, the same Chapter Header Block, Title Page,
# running heads, outer-edge page numbers, date-line style, drop cap on the first narration paragraph). The
# helpers are imported from area_docx_creator.py and area_md_creator.py, so a typeset fix made there reaches
# this file with no second copy of the rules.
#
# Length gate (author ruling 2026-10-07): the condensed Chapter holds between 50% and 60% of the characters
# (spaces counted) of the Chapter's written Areas. Both sides count the narration paragraphs only: the Area
# Title Block, the Title Page and Chapter Header Block, date lines, scene breaks and line breaks are left out.
# `write` refuses a text outside the band; `measure` reports it.
#
# TWO MODES (--mode, default condensed):
#   condensed  the original mode, for /chapter-builder: a condensed Chapter built from written Area files,
#              length gate 50%-60% of the Areas' characters. Unchanged.
#   whole      the mode for /create-chapter and /revise-chapter: a Chapter written whole, with no Area files.
#              Reads the Chapter Spec (read-only, --spec-file) for each Area block's `Size target`, `Budget stretch`
#              and number. The text carries one hidden line `<!-- AREA NN -->` at the start of every Area block of
#              the spec (in order, each exactly once). Word gate (words, not characters; date lines, markers and
#              breaks not counted): each Area block 55%-60% of its size target (an Area whose spec says
#              `Budget stretch: up to 72%` may reach 72%; none below 55%); Chapter total 0.55 T to 0.72 T,
#              T = the sum of the Area targets. `write` refuses an out-of-band text, a missing, duplicated, unknown or
#              out-of-order marker, a structure error, a Chapter that has Area files, and an existing file pair
#              without --overwrite. The .md keeps the hidden markers; the .docx carries none (stripped before
#              typesetting, no visible mark). `---` breaks are only where the text has them (the writer decides).
#
# Usage (command line; this is how /chapter-builder calls the script):
#   python .claude/tools/chapter_builder.py measure --text-file CH.txt --phase 1 --section 1 --part 1 --chapter 1
#   python .claude/tools/chapter_builder.py write --text-file CH.txt \
#       --phase 1 --section 1 --part 1 --chapter 1 \
#       --chapter-name "..." --phase-name "..." --section-name "..." --novel-part-name "..." \
#       [--first-of-phase] [--first-of-section] \
#       [--phase-epigraph-file P.txt] [--section-epigraph-file S.txt] [--output-root DIR] [--overwrite]
#   Whole mode: add  --mode whole --spec-file <Chapter Spec .md>  to measure or write (same other flags):
#   python .claude/tools/chapter_builder.py measure --mode whole --spec-file SPEC.md --text-file CH.txt \
#       --phase 1 --section 1 --part 1 --chapter 2
#   python .claude/tools/chapter_builder.py write --mode whole --spec-file SPEC.md --text-file CH.txt \
#       --phase 1 --section 1 --part 1 --chapter 2 --chapter-name "..." --phase-name "..." --section-name "..." \
#       --novel-part-name "..." [--first-of-phase] [--first-of-section] [--phase-epigraph-file P.txt] \
#       [--section-epigraph-file S.txt] [--output-root DIR] [--overwrite]
#   With no arguments the script runs its self-test.
# The Chapter text file is an input kept outside projects/ (the staging folder .claude/temp/), not a script and
# not a Chapter file.
# ═══════════════════════════════════════════════════════════════════════════

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.exceptions import PackageNotFoundError

import area_docx_creator as dx
import area_md_creator as mdc

DEFAULT_OUTPUT_ROOT = dx.DEFAULT_OUTPUT_ROOT
MIN_RATIO = 0.50
MAX_RATIO = 0.60
SCENE_BREAK = "---"
SCENE_BREAK_DOCX = "*   *   *"
_AREA_MD_RE = re.compile(r"^area-(\d{2})\.md$")
_ILLEGAL_NAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


# ─── PATHS ────────────────────────────────────────────────────────────────
def chapter_stem(chapter: int, chapter_name: str) -> str:
    """'Chapter 001 - The Capital Waking': the file name without its extension."""
    name = _ILLEGAL_NAME_CHARS.sub("", chapter_name or "").strip()
    return f"Chapter {chapter:03d}" + (f" - {name}" if name else "")


def get_part_dir(phase: int, section: int, part: int, output_root: str) -> str:
    return os.path.join(output_root, f"phase-{phase}", f"section-{section}", f"part-{part}")


def get_chapter_paths(phase: int, section: int, part: int, chapter: int,
                      chapter_name: str, output_root: str = DEFAULT_OUTPUT_ROOT) -> tuple:
    """(md path, docx path) of the Chapter file pair, in the Novel-Part folder."""
    base = os.path.join(get_part_dir(phase, section, part, output_root),
                        chapter_stem(chapter, chapter_name))
    return base + ".md", base + ".docx"


# ─── LENGTH GATE ──────────────────────────────────────────────────────────
def _body_paragraphs(text: str) -> list:
    """Narration paragraphs of a text: no blanks, scene breaks, headings, comments or date lines."""
    out = []
    for para in text.replace("\r\n", "\n").replace("\r", "\n").split("\n\n"):
        para = para.strip()
        if (not para or para in (SCENE_BREAK, "* * *", "***") or para.startswith(("#", "<!--"))
                or dx._date_line_text(para) is not None):
            continue
        out.append(para)
    return out


def _count(paragraphs: list) -> int:
    """Characters, spaces counted, of the paragraphs: line breaks joined as a space, emphasis marks left out."""
    return sum(len(p.replace("\n", " ").replace("*", "")) for p in paragraphs)


def area_files(phase: int, section: int, part: int, chapter: int, output_root: str) -> list:
    """(area number, path) of the written Area .md files of the Chapter, in Area order; _archive is never read."""
    folder = os.path.join(get_part_dir(phase, section, part, output_root), f"chapter-{chapter:03d}")
    return dx._sorted_numbered(folder, r"^area-(\d{2})\.md$")


def source_length(phase: int, section: int, part: int, chapter: int, output_root: str) -> int:
    """Characters of the Chapter's Areas: the narration after each Area's Reference line."""
    files = area_files(phase, section, part, chapter, output_root)
    if not files:
        raise FileNotFoundError("no written Area .md files found for the Chapter; nothing to condense")
    numbers = [n for n, _p in files]
    if numbers != list(range(1, len(numbers) + 1)):
        raise ValueError(f"the Area files found are {numbers}; they should run from 1 with no gap")
    total = 0
    for _n, path in files:
        text = dx._read_utf8(path)
        marker = re.search(r"^\*Reference:[^\n]*\*\s*$", text, flags=re.M)
        if not marker:
            raise ValueError(f"{path} has no *Reference:* line (the Area Title Block)")
        total += _count(_body_paragraphs(text[marker.end():]))
    return total


def measure(text: str, phase: int, section: int, part: int, chapter: int,
            output_root: str = DEFAULT_OUTPUT_ROOT) -> tuple:
    """(condensed chars, source chars, ratio, in band)."""
    condensed = _count(_body_paragraphs(text))
    source = source_length(phase, section, part, chapter, output_root)
    ratio = condensed / source if source else 0.0
    return condensed, source, ratio, MIN_RATIO <= ratio <= MAX_RATIO


# ─── TEXT STRUCTURE ───────────────────────────────────────────────────────
def _paragraphs(text: str) -> list:
    return [p.strip() for p in text.replace("\r\n", "\n").replace("\r", "\n").strip().split("\n\n") if p.strip()]


def _check_text(paras: list) -> None:
    if not paras or dx._date_line_text(paras[0]) is None:
        raise ValueError("the Chapter text must begin with its date line, e.g. *Year 1247, Ilovany, Orvil, Day 1*")
    if paras[-1] == SCENE_BREAK or paras[1:2] == [SCENE_BREAK]:
        raise ValueError("a scene break cannot open or close the Chapter text")
    for i, p in enumerate(paras):
        if p == SCENE_BREAK and (i + 1 >= len(paras) or paras[i + 1] == SCENE_BREAK):
            raise ValueError("a scene break must be followed by a scene")
    # a date line is allowed only first, or directly after a scene break
    for i, p in enumerate(paras):
        if i and dx._date_line_text(p) is not None and paras[i - 1] != SCENE_BREAK:
            raise ValueError("a date line may appear only at the top or directly after a scene break: " + p[:60])


# ─── DOCX ─────────────────────────────────────────────────────────────────
def _append_body_docx(doc, paras: list, font_name: str = dx.FONT_NAME) -> None:
    after_break = False
    opening = True      # the next narration paragraph takes the drop cap (after a date line)
    for p in paras:
        if p == SCENE_BREAK:
            dx._add_paragraph(doc, SCENE_BREAK_DOCX, size=dx.BODY_SIZE, align=WD_ALIGN_PARAGRAPH.CENTER,
                              first_indent=0, space_before=12, space_after=12, font_name=font_name)
            after_break = True
            continue
        if p in dx.TIME_SKIP_MARKERS:      # Rule 6.3.5 time-skip mark inside an Area: the same asterism the Area writer prints
            dx._add_paragraph(doc, "   ⁂", size=dx.ORNAMENT_SIZE, align=WD_ALIGN_PARAGRAPH.CENTER,
                              first_indent=0, space_before=12, space_after=12, font_name=font_name)
            after_break = False
            continue
        date = dx._date_line_text(p)
        if date is not None:
            dx._add_date_line_paragraph(doc, date, font_name=font_name)
            opening = True
            after_break = False
            continue
        segments = dx._parse_inline_formatting(p)
        if opening:
            dx._add_drop_cap_paragraph(doc, segments, font_name=font_name)
            opening = False
        else:
            dx._add_inline_formatted_paragraph(doc, segments,
                                               first_indent=0 if after_break else dx.FIRST_INDENT,
                                               font_name=font_name)
        after_break = False


def build_docx(paras, phase, section, part, chapter, phase_name, section_name, novel_part_name,
               chapter_name, first_of_phase, first_of_section, phase_epigraph, section_epigraph,
               output_root):
    doc = Document()
    dx._apply_page_setup(doc)
    dx._apply_running_headers_footers(doc, chapter_name, dx.FONT_NAME)
    if first_of_phase or first_of_section:
        dx._write_title_page(doc, phase, section, phase_name, section_name,
                             section_epigraph, output_root, phase_epigraph=phase_epigraph)
    dx._write_chapter_header_block(doc, phase, section, part, chapter, phase_name, section_name,
                                   novel_part_name, chapter_name, dx.get_phase_icon_path(phase, output_root))
    _append_body_docx(doc, paras)
    return doc


# ─── MD ───────────────────────────────────────────────────────────────────
def build_md(paras, phase, section, part, chapter, phase_name, section_name, novel_part_name,
             chapter_name, first_of_phase, first_of_section, phase_epigraph, section_epigraph) -> str:
    lines = []
    if first_of_phase or first_of_section:
        lines += ["---", "", f"# PHASE {phase}" + (f" — {phase_name.upper()}" if phase_name else ""), "",
                  f"## Section {section}" + (f" — {section_name}" if section_name else ""), ""]
        for line in (phase_epigraph, section_epigraph):
            if line and line.strip():
                lines += [f"*{line.strip()}*", ""]
        lines += ["---", ""]
    lines += mdc._md_chapter_header_block(phase, section, part, chapter, phase_name, section_name,
                                          novel_part_name, chapter_name)
    lines.append("")
    for p in paras:
        lines += [p, ""]
    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines) + "\n"


# ─── PRIMARY ENTRY POINT ──────────────────────────────────────────────────
def write_chapter(phase: int, section: int, part: int, chapter: int, text: str, chapter_name: str,
                  phase_name: str = "", section_name: str = "", novel_part_name: str = "",
                  first_of_phase: bool = False, first_of_section: bool = False,
                  phase_epigraph: str = "", section_epigraph: str = "",
                  output_root: str = DEFAULT_OUTPUT_ROOT, overwrite: bool = False) -> tuple:
    """
    Write the Chapter .md and .docx as a coupled pair into the Novel-Part folder. Returns (md path, docx path).
    Refuses, before any file is written: an empty chapter name; a Chapter text that breaks the structure above;
    a text outside the 50%-60% length band; a Phase epigraph given for a Chapter that does not open its Phase;
    a missing Phase or Section image when a Title Page is due; and an existing file unless overwrite is True
    (the command asks the author before it passes overwrite).
    """
    if not chapter_name:
        raise ValueError("a Chapter name is mandatory (Part 28 Area 18)")
    if phase_epigraph.strip() and not first_of_phase:
        raise ValueError("a Phase epigraph line is printed only with the first Chapter of a Phase")
    paras = _paragraphs(text)
    _check_text(paras)
    condensed, source, ratio, ok = measure(text, phase, section, part, chapter, output_root)
    if not ok:
        raise ValueError(f"the condensed Chapter holds {condensed:,} of {source:,} characters "
                         f"({ratio:.1%}); it must hold {MIN_RATIO:.0%} to {MAX_RATIO:.0%}. Nothing was written.")
    md_path, docx_path = get_chapter_paths(phase, section, part, chapter, chapter_name, output_root)
    for path in (md_path, docx_path):
        if os.path.exists(path) and not overwrite:
            raise FileExistsError(f"{path} already exists; it is replaced only after the author confirms "
                                  f"(--overwrite)")
    doc = build_docx(paras, phase, section, part, chapter, phase_name, section_name, novel_part_name,
                     chapter_name, first_of_phase, first_of_section, phase_epigraph, section_epigraph,
                     output_root)
    md = build_md(paras, phase, section, part, chapter, phase_name, section_name, novel_part_name,
                  chapter_name, first_of_phase, first_of_section, phase_epigraph, section_epigraph)
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    dx._save_joined(doc, docx_path)                     # the .docx is written whole or not at all
    with open(md_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(md)
    return md_path, docx_path


# ─── WHOLE MODE (/create-chapter, /revise-chapter) ────────────────────────
WHOLE_MIN = 0.55
WHOLE_MAX = 0.60
WHOLE_STRETCH_MAX = 0.72
_EPS = 1e-9
_MARKER_RE = re.compile(r"^\s*<!--\s*AREA\s+(\d+)\s*-->\s*$", re.I)
_AREA_HEAD_RE = re.compile(r"^###\s+Area\s+(\d+)\b", re.I)
_HEAD_RE = re.compile(r"^#{1,3}\s")
_TARGET_RE = re.compile(r"Size target:?\*{0,2}:?\s*(\d[\d,\.]*)", re.I)
_STRETCH_RE = re.compile(r"Budget stretch:?\*{0,2}:?\s*(.+)", re.I)


def parse_spec_areas(spec_text: str) -> list:
    """[(area number, size target in words, stretch named)] of the Chapter Spec's Area blocks, in order."""
    lines = spec_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks, cur = [], None
    for line in lines:
        m = _AREA_HEAD_RE.match(line)
        if m:
            cur = {"n": int(m.group(1)), "target": None, "stretch": False}
            blocks.append(cur)
            continue
        if _HEAD_RE.match(line):
            cur = None
            continue
        if cur is None:
            continue
        t = _TARGET_RE.search(line)
        if t and cur["target"] is None:
            cur["target"] = int(float(t.group(1).replace(",", "").rstrip(".")))
        st = _STRETCH_RE.search(line)
        if st:
            cur["stretch"] = st.group(1).strip().lower().lstrip("*_ ").startswith("up to")
    if not blocks:
        raise ValueError("the Chapter Spec has no '### Area NN' blocks")
    for b in blocks:
        if not b["target"]:
            raise ValueError(f"Area {b['n']:02d} of the Chapter Spec has no '- **Size target:** N words' line")
    nums = [b["n"] for b in blocks]
    if len(set(nums)) != len(nums):
        raise ValueError(f"the Chapter Spec repeats an Area block number: {nums}")
    return [(b["n"], b["target"], b["stretch"]) for b in blocks]


def _tokens(text: str) -> list:
    """Paragraph tokens ('marker', n) or ('para', text); a marker line is its own token wherever it sits."""
    out = []
    for para in text.replace("\r\n", "\n").replace("\r", "\n").strip().split("\n\n"):
        buf = []
        for line in para.split("\n"):
            m = _MARKER_RE.match(line)
            if m:
                if "\n".join(buf).strip():
                    out.append(("para", "\n".join(buf).strip()))
                buf = []
                out.append(("marker", int(m.group(1))))
            else:
                buf.append(line)
        if "\n".join(buf).strip():
            out.append(("para", "\n".join(buf).strip()))
    return out


def _words(para: str) -> int:
    return sum(1 for w in para.replace("*", "").split() if re.search(r"\w", w))


def area_file_names(phase: int, section: int, part: int, chapter: int, output_root: str) -> list:
    """Names of Area files (area-NN.md or .docx) in the Chapter's per-Area folder; a whole-built Chapter has none."""
    folder = os.path.join(get_part_dir(phase, section, part, output_root), f"chapter-{chapter:03d}")
    if not os.path.isdir(folder):
        return []
    return sorted(n for n in os.listdir(folder) if re.match(r"^area-\d{2}\.(md|docx)$", n))


def measure_whole(text: str, spec_text: str, phase: int = 0, section: int = 0, part: int = 0, chapter: int = 0,
                  output_root: str = DEFAULT_OUTPUT_ROOT) -> dict:
    """Every whole-mode check. Returns a dict; result['ok'] is True only when all of them pass."""
    spec = parse_spec_areas(spec_text)
    spec_nums = [n for n, _t, _s in spec]
    toks = _tokens(text)
    markers = [v for k, v in toks if k == "marker"]
    errors = []
    if markers != spec_nums:
        for n in spec_nums:
            c = markers.count(n)
            if c == 0:
                errors.append(f"marker <!-- AREA {n:02d} --> is missing")
            elif c > 1:
                errors.append(f"marker <!-- AREA {n:02d} --> appears {c} times")
        for n in sorted(set(markers) - set(spec_nums)):
            errors.append(f"marker <!-- AREA {n:02d} --> is not an Area block of the Chapter Spec")
        if not errors:
            errors.append(f"markers are out of order: found {markers}, the Chapter Spec runs {spec_nums}")
    counts = {n: 0 for n in spec_nums}
    cur, stray = None, 0
    for kind, v in toks:
        if kind == "marker":
            cur = v
            continue
        if v == SCENE_BREAK or dx._date_line_text(v) is not None or v.startswith("#"):
            continue
        w = _words(v)
        if cur is None:
            stray += w
        elif cur in counts:
            counts[cur] += w
    if stray:
        errors.append(f"{stray} words sit before the first Area marker")
    rows = []
    for n, target, stretch in spec:
        words = counts[n]
        ratio = words / target
        top = WHOLE_STRETCH_MAX if stretch else WHOLE_MAX
        ok = WHOLE_MIN - _EPS <= ratio <= top + _EPS
        rows.append({"n": n, "target": target, "words": words, "ratio": ratio, "stretch": stretch, "top": top,
                     "ok": ok})
    T = sum(r["target"] for r in rows)
    total = sum(r["words"] for r in rows)
    total_ok = WHOLE_MIN * T - _EPS <= total <= WHOLE_STRETCH_MAX * T + _EPS
    structure = None
    try:
        _check_text([v for k, v in toks if k == "para"])
    except ValueError as exc:
        structure = str(exc)
    areas = area_file_names(phase, section, part, chapter, output_root) if phase else []
    ok = not errors and all(r["ok"] for r in rows) and total_ok and structure is None and not areas
    return {"rows": rows, "T": T, "total": total, "ratio": total / T, "total_ok": total_ok,
            "marker_errors": errors, "structure_error": structure, "area_files": areas, "ok": ok}


def format_whole(res: dict) -> str:
    out = [f"{'Area':<6}{'target':>8}{'words':>8}{'ratio':>8}  band"]
    for r in res["rows"]:
        band = f"{WHOLE_MIN:.0%}-{r['top']:.0%}" + (" (stretch)" if r["stretch"] else "")
        out.append(f"{r['n']:02d}    {r['target']:>8,}{r['words']:>8,}{r['ratio']:>8.1%}  {band}  "
                   + ("PASS" if r["ok"] else "OUT OF BAND"))
    out.append(f"Chapter: {res['total']:,} words of T = {res['T']:,} ({res['ratio']:.1%}); allowed "
               f"{WHOLE_MIN * res['T']:,.0f} to {WHOLE_STRETCH_MAX * res['T']:,.0f} words: "
               + ("PASS" if res["total_ok"] else "OUT OF BAND"))
    for e in res["marker_errors"]:
        out.append("markers: " + e)
    if res["structure_error"]:
        out.append("structure: " + res["structure_error"])
    if res["area_files"]:
        out.append("area files exist for the Chapter (a whole-built Chapter has none): "
                   + ", ".join(res["area_files"]))
    out.append("result: " + ("PASS" if res["ok"] else "FAIL"))
    return "\n".join(out)


def write_chapter_whole(phase: int, section: int, part: int, chapter: int, text: str, spec_text: str,
                        chapter_name: str, phase_name: str = "", section_name: str = "", novel_part_name: str = "",
                        first_of_phase: bool = False, first_of_section: bool = False,
                        phase_epigraph: str = "", section_epigraph: str = "",
                        output_root: str = DEFAULT_OUTPUT_ROOT, overwrite: bool = False) -> tuple:
    """
    Write a whole-built Chapter pair. Refuses, before any file is written, everything measure_whole reports as a
    failure (band, markers, structure, Area files) plus the checks of write_chapter (name, epigraph, existing files).
    The .md keeps the hidden Area markers; the .docx has them stripped. Returns (md path, docx path).
    """
    if not chapter_name:
        raise ValueError("a Chapter name is mandatory (Part 28 Area 18)")
    if phase_epigraph.strip() and not first_of_phase:
        raise ValueError("a Phase epigraph line is printed only with the first Chapter of a Phase")
    res = measure_whole(text, spec_text, phase, section, part, chapter, output_root)
    if not res["ok"]:
        raise ValueError("the Chapter text fails the whole-mode checks. Nothing was written.\n" + format_whole(res))
    toks = _tokens(text)
    md_paras = [("<!-- AREA %02d -->" % v) if k == "marker" else v for k, v in toks]
    docx_paras = [v for k, v in toks if k == "para"]
    md_path, docx_path = get_chapter_paths(phase, section, part, chapter, chapter_name, output_root)
    for path in (md_path, docx_path):
        if os.path.exists(path) and not overwrite:
            raise FileExistsError(f"{path} already exists; it is replaced only after the author confirms "
                                  f"(--overwrite)")
    doc = build_docx(docx_paras, phase, section, part, chapter, phase_name, section_name, novel_part_name,
                     chapter_name, first_of_phase, first_of_section, phase_epigraph, section_epigraph,
                     output_root)
    md = build_md(md_paras, phase, section, part, chapter, phase_name, section_name, novel_part_name,
                  chapter_name, first_of_phase, first_of_section, phase_epigraph, section_epigraph)
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    dx._save_joined(doc, docx_path)
    with open(md_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(md)
    return md_path, docx_path


# ─── COMMAND LINE ─────────────────────────────────────────────────────────
def _build_parser():
    parser = argparse.ArgumentParser(
        prog="chapter_builder.py", allow_abbrev=False,
        description="Starward Saga Chapter writer, condensed mode (default) and whole mode (authorized script 6, "
                    "guardrails/tools/guardrail-authorized-scripts-only.md). With no arguments it runs the self-test.")
    sub = parser.add_subparsers(dest="command", required=True)
    for name, help_ in (("measure", "report the text's length (condensed: against the Areas; whole: word bands per Area block)"),
                        ("write", "write the Chapter .md and .docx into the Novel-Part folder")):
        s = sub.add_parser(name, allow_abbrev=False, help=help_)
        s.add_argument("--mode", choices=("condensed", "whole"), default="condensed",
                       help="condensed (default, /chapter-builder: condensed from written Area files) or whole "
                            "(/create-chapter: a Chapter written whole, word bands checked against --spec-file)")
        s.add_argument("--spec-file", default="",
                       help="whole mode: the Chapter Spec .md (read-only) for Area size targets and stretches")
        s.add_argument("--text-file", required=True,
                       help="UTF-8 file holding the confirmed Chapter text (whole mode: with <!-- AREA NN --> markers)")
        for field in ("phase", "section", "part", "chapter"):
            s.add_argument("--" + field, type=int, required=True)
        s.add_argument("--output-root", default=DEFAULT_OUTPUT_ROOT)
        if name == "write":
            for field in ("chapter-name", "phase-name", "section-name", "novel-part-name"):
                s.add_argument("--" + field, default="")
            s.add_argument("--first-of-phase", action="store_true")
            s.add_argument("--first-of-section", action="store_true")
            s.add_argument("--phase-epigraph", default="")
            s.add_argument("--phase-epigraph-file", default="")
            s.add_argument("--section-epigraph", default="")
            s.add_argument("--section-epigraph-file", default="")
            s.add_argument("--overwrite", action="store_true",
                           help="replace an existing Chapter file (after the author has confirmed)")
    return parser


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    argv = sys.argv[1:] if argv is None else list(argv)
    if not argv:
        _self_test()
        return 0
    parser = _build_parser()
    args = parser.parse_args(argv)
    if args.mode == "whole" and not args.spec_file:
        parser.error("--mode whole needs --spec-file <Chapter Spec path>")
    root = os.path.abspath(args.output_root)
    try:
        text = dx._read_utf8(args.text_file)
        spec_text = dx._read_utf8(args.spec_file) if args.mode == "whole" else ""
        if args.command == "measure" and args.mode == "whole":
            res = measure_whole(text, spec_text, args.phase, args.section, args.part, args.chapter, root)
            print(format_whole(res))
            return 0 if res["ok"] else 1
        if args.command == "measure":
            condensed, source, ratio, ok = measure(text, args.phase, args.section, args.part, args.chapter, root)
            print(f"condensed: {condensed:,} characters (spaces counted)")
            print(f"Areas:     {source:,} characters")
            print(f"ratio:     {ratio:.1%} (band {MIN_RATIO:.0%} to {MAX_RATIO:.0%})")
            print("result: " + ("IN BAND" if ok else "OUT OF BAND"))
            return 0 if ok else 1

        def _given(inline, file_):
            value = (dx._read_utf8(file_) if file_ else inline).strip()
            return "" if value.lower() == "unset" else value
        extra = (spec_text,) if args.mode == "whole" else ()
        md_path, docx_path = (write_chapter_whole if args.mode == "whole" else write_chapter)(
            args.phase, args.section, args.part, args.chapter, text, *extra, args.chapter_name,
            phase_name=args.phase_name, section_name=args.section_name, novel_part_name=args.novel_part_name,
            first_of_phase=args.first_of_phase, first_of_section=args.first_of_section,
            phase_epigraph=_given(args.phase_epigraph, args.phase_epigraph_file),
            section_epigraph=_given(args.section_epigraph, args.section_epigraph_file),
            output_root=root, overwrite=args.overwrite)
        print("wrote: " + os.path.abspath(md_path))
        print("wrote: " + os.path.abspath(docx_path))
        return 0
    except (FileExistsError, FileNotFoundError, ValueError, PackageNotFoundError) as exc:
        print("refused: " + str(exc), file=sys.stderr)
        return 2


# ─── SELF-TEST ────────────────────────────────────────────────────────────
def _self_test() -> None:
    import tempfile
    print("Starward Saga Chapter Builder — authorized script 6")
    with tempfile.TemporaryDirectory() as tmp:
        cdir = os.path.join(get_part_dir(1, 1, 1, tmp), "chapter-001")
        os.makedirs(cdir)
        for n in (1, 2):
            with open(os.path.join(cdir, f"area-{n:02d}.md"), "w", encoding="utf-8") as fh:
                fh.write(f"## Area {n} — X\n*Reference: 1.1.1.001.0{n}*\n\n*Year 1247, A, B, Day 1*\n\n"
                         + ("word " * 100).strip() + "\n\n" + ("more " * 100).strip() + "\n")
        dl = "*Year 1247, A, B, Day 1*"
        text = f"{dl}\n\n" + ("word " * 115).strip() + "\n\n---\n\n" + ("tale " * 115).strip()
        text2 = f"{dl}\n\n" + ("word " * 115).strip() + "\n\n---\n\n*Year 1247, A, B, Day 2*\n\n" + ("tale " * 115).strip()
        short = f"{dl}\n\nshort."
        refused = False
        try:
            write_chapter(1, 1, 1, 1, short, "C", output_root=tmp)
        except ValueError:
            refused = True
        md_path, docx_path = write_chapter(1, 1, 1, 1, text2, "The Test", "P", "S", "N", output_root=tmp)
        ps = Document(docx_path).paragraphs
        texts = [p.text for p in ps]
        i = texts.index("Year 1247, A, B, Day 2")
        d1 = texts.index("Year 1247, A, B, Day 1")
        ok = (refused and md_path.replace("\\", "/").endswith("part-1/Chapter 001 - The Test.md")
              and docx_path.replace("\\", "/").endswith("part-1/Chapter 001 - The Test.docx")
              and texts.count(SCENE_BREAK_DOCX) == 1
              and ps[d1 + 1].runs[0].font.size.pt == dx.DROP_CAP_SIZE
              and ps[i + 1].runs[0].font.size.pt == dx.DROP_CAP_SIZE
              and not any(t.startswith("Area ") for t in texts)
              and "---" in open(md_path, encoding="utf-8").read()
              and measure(text, 1, 1, 1, 1, tmp)[3])
        print("Chapter file sample test:", "PASS" if ok else "FAIL")
        if not ok:
            sys.exit(1)
    _self_test_whole()


def _self_test_whole() -> None:
    import tempfile
    spec = ("# Chapter\n**Chapter size target (words):** 1,000\n**Build mode:** whole\n\n"
            "### Area 01 — A\n**Pacing (planned):** Slow\n- **Size target:** 400 words\n- **Budget stretch:** none\n\n"
            "### Area 02 — B\n**Pacing (planned):** Fast\n- **Size target:** 400 words\n"
            "- **Budget stretch:** up to 72% (the fight runs long)\n\n"
            "### Area 03 — C\n**Pacing (planned):** Medium\n- **Size target:** 200 words\n- **Budget stretch:** none\n")

    def words(n, w="word"):
        return " ".join([w] * n)
    d1, d2 = "*Year 1247, A, B, Day 1*", "*Year 1247, A, B, Day 2*"
    # A1 = 220 (55%), A2 = 280 (70%, stretch), A3 = 120 (60%); total 620 of T=1000 (62%)
    good = (f"<!-- AREA 01 -->\n{d1}\n\n{words(110)}\n\n{words(110, 'tale')}\n\n<!-- AREA 02 -->\n{words(140, 'more')}\n\n"
            f"---\n\n{d2}\n\n{words(140, 'fight')}\n\n<!-- AREA 03 -->\n{words(120, 'end')}\n")
    with tempfile.TemporaryDirectory() as tmp:
        def refused(**kw):
            try:
                write_chapter_whole(1, 1, 1, 2, kw.pop("text", good), spec, "The Whole", output_root=tmp, **kw)
            except (ValueError, FileExistsError):
                return True
            return False
        res = measure_whole(good, spec, 1, 1, 1, 2, tmp)
        no_stretch_spec = spec.replace("up to 72% (the fight runs long)", "none")
        stretch_parsed = [r["stretch"] for r in res["rows"]] == [False, True, False]
        out_of_band = measure_whole(good, no_stretch_spec, 1, 1, 1, 2, tmp)["ok"] is False   # 70% needs the stretch
        r_band = refused(text=good.replace(words(120, "end"), words(40, "end")))             # Area 03 at 20%
        r_missing = refused(text=good.replace("<!-- AREA 03 -->\n", ""))
        r_dup = refused(text=good + "\n<!-- AREA 03 -->\n" + words(5))
        r_order = refused(text=good.replace("<!-- AREA 02 -->", "<!-- AREA 09 -->"))
        r_open = refused(text=good.replace(d1 + "\n\n", "", 1))                             # opening date line gone
        r_total = refused(text=good.replace(words(140, "fight"), words(250, "fight")))      # Area 02 over 72%
        os.makedirs(os.path.join(get_part_dir(1, 1, 1, tmp), "chapter-002"))
        open(os.path.join(get_part_dir(1, 1, 1, tmp), "chapter-002", "area-01.md"), "w").close()
        r_areas = refused()
        os.remove(os.path.join(get_part_dir(1, 1, 1, tmp), "chapter-002", "area-01.md"))
        md_path, docx_path = write_chapter_whole(1, 1, 1, 2, good, spec, "The Whole", "P", "S", "N", output_root=tmp)
        r_exists = refused()
        md = open(md_path, encoding="utf-8").read()
        texts = [p.text for p in Document(docx_path).paragraphs]
        ps = Document(docx_path).paragraphs
        i2 = texts.index("Year 1247, A, B, Day 2")
        ok = (res["ok"] and stretch_parsed and out_of_band and r_band and r_missing and r_dup and r_order
              and r_open and r_total and r_areas and r_exists
              and all(f"<!-- AREA {n:02d} -->" in md for n in (1, 2, 3))
              and not any("AREA" in t for t in texts)
              and texts.count(SCENE_BREAK_DOCX) == 1
              and ps[i2 + 1].runs[0].font.size.pt == dx.DROP_CAP_SIZE)
        print("Whole-mode sample test:", "PASS" if ok else "FAIL")
        if not ok:
            print(format_whole(res))
            print(dict(stretch_parsed=stretch_parsed, out_of_band=out_of_band, r_band=r_band, r_missing=r_missing,
                       r_dup=r_dup, r_order=r_order, r_open=r_open, r_total=r_total, r_areas=r_areas,
                       r_exists=r_exists))
            sys.exit(1)


if __name__ == "__main__":
    sys.exit(main())
